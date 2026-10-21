"""Package the unchanged supplied game with the QA handoff, never credentials."""
from pathlib import Path
import hashlib, json, tarfile, zipfile, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/home/user/webapp')
QA=ROOT/'qa'
ART=ROOT/'artifacts'
ART.mkdir(exist_ok=True)

# Compare every original archive file, not just tracked JS files.
checks=[]
with tarfile.open('/home/user/uploaded_files/LUMEN.tar.gz') as archive:
    for item in archive.getmembers():
        if not item.isfile():
            continue
        original=archive.extractfile(item).read()
        actual=(ROOT/item.name).read_bytes()
        checks.append({'path':item.name,'sha256':hashlib.sha256(actual).hexdigest(),'unchanged':original==actual})
assert all(item['unchanged'] for item in checks), 'Imported input was modified'
manifest={'input':'uploaded LUMEN.tar.gz; NOT verified GitHub HEAD','filesCompared':len(checks),'allUnchanged':True,'files':checks}
(QA/'evidence/archive-integrity.json').write_text(json.dumps(manifest,indent=2))

# Editable Word document from the authored Markdown.
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.75)
normal=doc.styles['Normal']
normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
for level in range(1,4):
    doc.styles[f'Heading {level}'].font.color.rgb=RGBColor.from_string('213B57')
for line in (QA/'LUMEN_BUG_REPORT.md').read_text().splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:],0)
    elif line.startswith('## '):
        doc.add_heading(line[3:],1)
    elif line.startswith('### '):
        doc.add_heading(line[4:],2)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:],style='List Bullet')
    else:
        doc.add_paragraph(line)
footer=sec.footer.paragraphs[0]
footer.text='LUMEN QA handoff | Report only, game unchanged | Page '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='LUMEN LEDGER Beta-test bug handoff'
doc.core_properties.subject='Two-run beta test, defects, evidence, and next-chat handoff'
doc.save(QA/'LUMEN_BUG_REPORT.docx')

zip_path=ART/'LUMEN_game_and_QA.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED,compresslevel=8) as out:
    for file in sorted((ROOT/'LUMEN/game').rglob('*')):
        if file.is_file():
            out.write(file,file.relative_to(ROOT/'LUMEN/game'))
    for file in sorted(QA.rglob('*')):
        if file.is_file() and file.suffix not in ('.pyc',):
            out.write(file,Path('QA')/file.relative_to(QA))
with zipfile.ZipFile(zip_path) as test:
    assert test.testzip() is None
    assert 'index.html' in test.namelist()
    assert 'QA/LUMEN_BUG_REPORT.docx' in test.namelist()
    assert all(not name.startswith('/') and '..' not in Path(name).parts for name in test.namelist())
    for name in test.namelist():
        if name.endswith(('.js','.mjs','.md','.txt','.json','.py','.html')):
            assert not re.search(rb'gh[pousr]_[A-Za-z0-9]{20,}',test.read(name)), 'Credential-like token in delivery'
print(json.dumps({'zip':str(zip_path),'zipBytes':zip_path.stat().st_size,'zipSha256':hashlib.sha256(zip_path.read_bytes()).hexdigest(),'originalFilesCompared':len(checks),'gameUnchanged':True,'docx':str(QA/'LUMEN_BUG_REPORT.docx')},indent=2))
